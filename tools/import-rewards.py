"""Read verified reward milestones and task-reference ranges into rewards.json.

Run with the workbook path as an argument. The workbook is opened read-only and is
never saved. The independent generic "每日任務" block remains outside both event modes.
"""

import json
import re
import sys
import struct
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


CURRENCY = {
    "GOLD": "gold",
    "NODE_STONE": "core",
    "DICETREE_REFUND": "treeSeed",
    "TICKET_COOP": "coopTicket",
    "TICKET_ARENA": "arenaTicket",
    "SKIN_COIN": "skinCoin",
}
REPEATABLE_CURRENCY = {**CURRENCY, "TICKET_SIXDICE": "luckyDiceTicket"}
DICE_NODES = {
    "UNIQUE:Dice_BINGO3": "4008",
    "RAW:Dice_Slow3": "5006",
    "RAW:Dice_SPEEDGUN3": "5008",
}
# User-confirmed character labels; source workbook stays untouched.
AVATAR_LABELS = {
    "自然系頭像": "青兒頭像",
    "秩序系頭像": "迪奇頭像", "入侵系頭像": "迪奇頭像",
    "工學系頭像": "里克頭像", "工程系頭像": "里克頭像",
    "魔法系頭像": "艾科頭像", "渾沌系頭像": "伊安頭像",
}
MODE_RANGES = [
    ("raid-normal", "討伐一般", "累計擊殺", "模式獎勵系統", 5, 35),
    ("raid-hard", "討伐困難", "累計擊殺", "模式獎勵系統", 39, 88),
    ("arena-pass", "競技場通行證", "累計積分", "模式獎勵系統", 92, 110),
    ("journey-7day", "7日旅程", "累計積分", "任務積分系統", 42, 46),
    ("hunt-event", "狩獵活動", "累計積分", "任務積分系統", 122, 135),
]
TASK_RANGES = {
    "journey-7day": (5, 39, 7),
    "hunt-event": (50, 119, 14),
}


def integer(value):
    if isinstance(value, str):
        value = value.replace(",", "")
    number = int(value)
    assert float(value) == number and number >= 0, value
    return number


def main():
    if len(sys.argv) != 2:
        raise SystemExit("usage: import-rewards.py <master-workbook.xlsx>")
    workbook = load_workbook(sys.argv[1], read_only=True, data_only=True)
    fixed = workbook["固定獎勵統計"]
    # The workbook owns the classification. This map only joins its official
    # reward IDs to the existing website currency icon registry keys.
    fixed_currency_ids = {fixed.cell(row, 10).value for row in range(5, 11)}
    assert fixed_currency_ids == set(CURRENCY), fixed_currency_ids ^ set(CURRENCY)
    labels = {fixed.cell(row, 10).value: fixed.cell(row, 5).value
              for row in range(14, 36) if fixed.cell(row, 10).value}
    assert len(labels) == 20

    def grant(item_id, amount):
        amount = integer(amount)
        assert amount > 0
        if item_id in fixed_currency_ids:
            return {"type": "currency", "kind": CURRENCY[item_id], "amount": amount}
        if item_id in DICE_NODES:
            return {"type": "dice", "itemId": item_id, "label": labels[item_id],
                    "nodeId": DICE_NODES[item_id], "assetStatus": "ready", "amount": amount}
        if item_id == "UNIQUE:target":
            return {"type": "emote", "itemId": item_id, "label": labels[item_id],
                    "assetStatus": "ready", "icon": "/rewards/emote/target.png", "amount": amount}
        if item_id == "UNIQUE:Dice_Predator3_skin1":
            return {"type": "cosmetic", "subtype": "dice-skin", "itemId": item_id,
                    "label": "吞噬骰子－鯊魚造型", "assetStatus": "ready",
                    "icon": "/rewards/cosmetic/Dice_Predator3_skin1.png",
                    "amount": amount}
        if isinstance(item_id, str) and item_id.startswith("RAW:") and item_id in labels:
            subtype = ("banner" if item_id.startswith("RAW:Profile_Banner_") else
                       "frame" if item_id.startswith("RAW:rd2_ui_profile_rank_border_") else "avatar")
            return {"type": "cosmetic", "subtype": subtype, "itemId": item_id,
                    "label": AVATAR_LABELS.get(labels[item_id], labels[item_id]),
                    "icon": f"/rewards/cosmetic/{item_id[4:]}.png", "assetStatus": "ready",
                    "amount": amount}
        raise ValueError(f"Unmapped reward ID: {item_id}")

    modes = []
    for mode_id, name, requirement_label, sheet_name, first, last in MODE_RANGES:
        sheet = workbook[sheet_name]
        tiers = []
        for row in range(first, last + 1):
            requirement = integer(sheet.cell(row, 1).value)
            assert sheet.cell(row, 17).value == "已確認", (sheet_name, row)
            rewards = []
            for id_col in (11, 13, 15):
                item_id = sheet.cell(row, id_col).value
                amount = sheet.cell(row, id_col + 1).value
                if item_id:
                    rewards.append(grant(item_id, amount))
                elif amount is not None:
                    # The master has no official ID for these two emotes. The
                    # source-cell key is an internal slot, not an invented game ID.
                    assert sheet_name == "模式獎勵系統" and row in (44, 78)
                    label = "STOP 表情" if row == 44 else "NO_SIGN 表情"
                    emote = "STOP" if row == 44 else "NO_SIGN"
                    rewards.append({"type": "emote", "itemId": f"source:{sheet_name}:{row}:{id_col}",
                                    "label": label, "assetStatus": "ready",
                                    "icon": f"/rewards/emote/{emote}.png", "amount": integer(amount)})
            assert rewards, (sheet_name, row)
            tiers.append({"id": f"{mode_id}-{row}", "requirement": requirement, "rewards": rewards})
        assert len(set(tier["requirement"] for tier in tiers)) == len(tiers)
        mode = {"kind": "threshold", "id": mode_id, "name": name,
                "requirementLabel": requirement_label, "tiers": tiers}
        if mode_id in TASK_RANGES:
            task_first, task_last, day_count = TASK_RANGES[mode_id]
            task_sheet = workbook["任務積分系統"]
            days = []
            for row in range(task_first, task_last + 1):
                day_text = task_sheet.cell(row, 1).value
                if day_text:
                    assert day_text == f"第{len(days) + 1}天", (row, day_text)
                    days.append({"day": len(days) + 1,
                                 "totalPoints": integer(task_sheet.cell(row, 7).value),
                                 "tasks": []})
                assert days and task_sheet.cell(row, 2).value, row
                days[-1]["tasks"].append({
                    "id": f"{mode_id}-task-{row}",
                    "name": str(task_sheet.cell(row, 2).value).strip(),
                    "requirement": integer(task_sheet.cell(row, 5).value),
                    "points": integer(task_sheet.cell(row, 6).value),
                })
            assert len(days) == day_count
            assert all(len(day["tasks"]) == 5 and
                       sum(task["points"] for task in day["tasks"]) == day["totalPoints"]
                       for day in days)
            mode["taskDays"] = days
        modes.append(mode)

    sheet = workbook["成就獎勵系統"]
    groups = {}
    current_name = None
    for row in range(4, 314):
        stage_text = sheet.cell(row, 2).value
        if not stage_text:
            continue
        assert sheet.cell(row, 11).value == "已確認", row
        if sheet.cell(row, 1).value:
            current_name = str(sheet.cell(row, 1).value).strip()
        assert current_name
        match = re.fullmatch(r"(\d+)\s*/\s*(\d+)", str(stage_text).strip())
        assert match, (row, stage_text)
        stage_number, total_stages = map(int, match.groups())
        group = groups.setdefault(current_name, {"id": current_name, "name": current_name, "stages": []})
        assert stage_number == len(group["stages"]) + 1, (row, current_name, stage_number)
        assert stage_number <= total_stages
        group["stages"].append({"id": f"achievement-{row}", "stage": stage_number,
                                "requirement": integer(sheet.cell(row, 3).value),
                                "rewards": [grant(sheet.cell(row, 9).value, sheet.cell(row, 10).value)]})
    assert len(groups) == 39 and sum(len(group["stages"]) for group in groups.values()) == 274
    modes.append({"kind": "achievement", "id": "achievements", "name": "成就",
                  "groups": list(groups.values())})

    counts = [sum(len(tier["rewards"]) for tier in mode["tiers"])
              if mode["kind"] == "threshold" else 274 for mode in modes]
    assert counts == [54, 135, 29, 13, 30, 274], counts
    types = Counter(reward["type"] for mode in modes for stage in
                    (mode["tiers"] if mode["kind"] == "threshold" else
                     [stage for group in mode["groups"] for stage in group["stages"]])
                    for reward in stage["rewards"])
    assert types == {"currency": 513, "dice": 3, "cosmetic": 16, "emote": 3}, types
    # Read each repeatable milestone's own source row; never derive from totals.
    def reference_section(sheet_name, first, last, requirement_col, section_id, title, metric):
        sheet = workbook[sheet_name]
        tiers = []
        for row in range(first, last + 1):
            rewards = []
            for id_col in (11, 13, 15):
                item_id = sheet.cell(row, id_col).value
                if item_id:
                    if item_id in REPEATABLE_CURRENCY:
                        reward = {"type": "currency", "kind": REPEATABLE_CURRENCY[item_id],
                                  "amount": integer(sheet.cell(row, id_col + 1).value)}
                    else:
                        reward = grant(item_id, sheet.cell(row, id_col + 1).value)
                    assert reward["amount"] > 0
                    rewards.append(reward)
            assert rewards, (sheet_name, row)
            tiers.append({"id": f"{section_id}-{row}",
                          "requirement": integer(sheet.cell(row, requirement_col).value), "rewards": rewards})
        return {"id": section_id, "title": title, "metricLabel": metric, "tiers": tiers}

    streak = reference_section("模式獎勵系統", 114, 119, 2, "arena-streak", "競技場連勝", "連勝場數")
    streak["note"] = "連勝中斷時依最終連勝數發放；各檔不是累加獎勵。"
    # User-confirmed mapping: N wins -> complete official N_pig.png.
    streak["tierImages"] = {}
    for tier in streak["tiers"]:
        wins = tier["requirement"]
        file = f"{wins}_pig.png"
        pig = Path(__file__).resolve().parents[1] / "public/rewards/reference" / file
        width, height = struct.unpack(">II", pig.read_bytes()[16:24])
        streak["tierImages"][str(wins)] = {"src": f"/rewards/reference/{file}", "alt": f"{wins} 連勝撲滿", "width": width, "height": height}
    assert {workbook["模式獎勵系統"].cell(row, 10).value for row in range(114, 120)} == {"連勝中斷時依最終連勝數發放"}
    task_sheet = workbook["任務積分系統"]
    daily_tasks = [{"id": f"daily-task-{row}", "name": str(task_sheet.cell(row, 1).value).strip(),
                    "requirement": integer(task_sheet.cell(row, 5).value),
                    "points": integer(task_sheet.cell(row, 6).value)} for row in range(139, 147)]
    total_points = integer(task_sheet.cell(139, 7).value)
    assert sum(task["points"] for task in daily_tasks) == total_points
    daily_rewards = reference_section("任務積分系統", 149, 153, 1, "daily-points", "每日累積獎勵", "累積點數")
    # Preserve the workbook's unresolved annotation instead of claiming it was confirmed.
    assert {task_sheet.cell(row, 17).value for row in range(149, 154)} == {"獎勵圖示待辨識"}
    daily_rewards["note"] = "主表此區仍註記「獎勵圖示待辨識」；內容依已填寫的資源 ID 與數量呈現。"
    daily_rewards["title"] = "獎勵"
    repeatable = {"id": "repeatable", "name": "重複性獎勵", "sections": [streak],
                  "dailyTasks": {"tasks": daily_tasks, "totalPoints": total_points, "rewards": daily_rewards}}
    result = {"status": "official", "note": "資料取自主表已確認的獎勵里程碑；缺圖收藏品以文字顯示。",
              "modes": modes, "repeatable": repeatable}
    workbook.close()
    output = Path(__file__).resolve().parent.parent / "data" / "rewards.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{output}: {counts}; {dict(types)}")


if __name__ == "__main__":
    main()
